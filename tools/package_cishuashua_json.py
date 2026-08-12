#!/usr/bin/env python3
"""Package Cishuashua word rows from Excel into per-word JSON files and a zip.

Rules:
- Use a sample JSON as the template.
- Preserve tagId, componentName, and componentKey from the sample.
- Fill attrs.topicId, attrs.word, attrs.meanId, attrs.mean.
- Fill attrs.styles.<style> fields from spreadsheet columns.
- videoId = topicId + meanId + style. The default style is "default".
- videoSrc is read from "压缩后视频链接" by default.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Any

import openpyxl


DEFAULT_SAMPLE_JSON = Path(
    "/Users/admin/Library/Containers/com.tencent.WeWorkMac/Data/Documents/Profiles/ED7EEBABA5C160DF4E737A372B0E7DEC/Caches/Files/2026-07/6eb5ab200bca89c529dfdcab0cab58fb/5575-312-669-3126691782900633780373135.json"
)

REQUIRED_COLUMNS = [
    "topicId",
    "word",
    "Meaning-each",
    "MeaningId",
    "例句",
    "例句音频",
    "例句翻译",
]


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def parse_int(value: Any, column: str, excel_row: int) -> int:
    text = clean(value)
    if not text:
        raise ValueError(f"Row {excel_row}: column {column} is empty")
    try:
        return int(float(text))
    except ValueError as exc:
        raise ValueError(f"Row {excel_row}: column {column} is not an integer: {text!r}") from exc


def safe_filename(word: str) -> str:
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", word).strip("_").lower()
    return base or "word"


def load_template(path: Path, style: str) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8"))
    try:
        default_style = data["attrs"]["styles"]["default"]
    except KeyError as exc:
        raise ValueError("Sample JSON must contain attrs.styles.default") from exc

    styles = data["attrs"].setdefault("styles", {})
    if style not in styles:
        styles[style] = deepcopy(default_style)
    return data


def get_headers(sheet: Any) -> dict[str, int]:
    headers = [clean(cell.value) for cell in sheet[1]]
    return {header: index for index, header in enumerate(headers) if header}


def iter_data_rows(sheet: Any, start_index: int, count: int | None) -> list[tuple[int, tuple[Any, ...]]]:
    min_row = start_index + 1
    max_row = sheet.max_row if count is None else min_row + count - 1
    rows = []
    for excel_row, row in enumerate(
        sheet.iter_rows(min_row=min_row, max_row=max_row, values_only=True),
        start=min_row,
    ):
        if any(clean(value) for value in row):
            rows.append((excel_row, row))
    return rows


def build_word_json(
    template: dict[str, Any],
    row: tuple[Any, ...],
    columns: dict[str, int],
    excel_row: int,
    style: str,
    video_column: str,
    poster_column: str | None,
) -> dict[str, Any]:
    data = deepcopy(template)

    topic_id = parse_int(row[columns["topicId"]], "topicId", excel_row)
    mean_id = parse_int(row[columns["MeaningId"]], "MeaningId", excel_row)
    word = clean(row[columns["word"]])
    if not word:
        raise ValueError(f"Row {excel_row}: column word is empty")

    attrs = data["attrs"]
    attrs["topicId"] = topic_id
    attrs["word"] = word
    attrs["meanId"] = mean_id
    attrs["mean"] = clean(row[columns["Meaning-each"]])

    styles = attrs.setdefault("styles", {})
    if style not in styles:
        styles[style] = deepcopy(styles.get("default", {}))
    target_style = styles[style]

    target_style["videoId"] = f"{topic_id}{mean_id}{style}"
    target_style["videoSrc"] = clean(row[columns[video_column]])
    target_style["videoPoster"] = clean(row[columns[poster_column]]) if poster_column else ""
    target_style["example"] = clean(row[columns["例句"]])
    target_style["exampleAudio"] = clean(row[columns["例句音频"]])
    target_style["exampleTranslation"] = clean(row[columns["例句翻译"]])

    return data


def make_package(args: argparse.Namespace) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(args.xlsx, data_only=True)
    sheet = workbook[args.sheet] if args.sheet else workbook[workbook.sheetnames[0]]

    columns = get_headers(sheet)
    missing = [column for column in REQUIRED_COLUMNS if column not in columns]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    video_column = args.video_column
    if video_column not in columns:
        raise ValueError(f"Missing video source column: {video_column}")

    poster_column = args.poster_column
    if poster_column and poster_column not in columns:
        poster_column = None

    template = load_template(args.sample_json, args.style)
    rows = iter_data_rows(sheet, args.start_index, args.count)

    if args.clean and args.out_dir.exists():
        shutil.rmtree(args.out_dir)
    args.out_dir.mkdir(parents=True, exist_ok=True)
    args.zip.parent.mkdir(parents=True, exist_ok=True)

    files = []
    used_names: dict[str, int] = {}
    errors = []

    for excel_row, row in rows:
        try:
            data = build_word_json(
                template=template,
                row=row,
                columns=columns,
                excel_row=excel_row,
                style=args.style,
                video_column=video_column,
                poster_column=poster_column,
            )
            word = data["attrs"]["word"]
            base = safe_filename(word)
            suffix = used_names.get(base, 0) + 1
            used_names[base] = suffix
            filename = f"{base}.json" if suffix == 1 else f"{base}_{suffix}.json"
            output_path = args.out_dir / filename
            output_path.write_text(json.dumps(data, ensure_ascii=False, indent=4), encoding="utf-8")
            files.append(output_path)
        except Exception as exc:  # noqa: BLE001 - return row-level validation errors to caller.
            errors.append({"row": excel_row, "error": str(exc)})
            if not args.continue_on_error:
                break

    if errors and not args.continue_on_error:
        raise ValueError(json.dumps({"errors": errors}, ensure_ascii=False))

    with zipfile.ZipFile(args.zip, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in files:
            archive.write(path, arcname=path.name)

    return {
        "ok": not errors,
        "zip": str(args.zip.resolve()),
        "out_dir": str(args.out_dir.resolve()),
        "count": len(files),
        "errors": errors,
        "sheet": sheet.title,
        "style": args.style,
        "video_column": video_column,
        "poster_column": poster_column,
        "files": [path.name for path in files],
    }


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Package Cishuashua Excel rows into JSON zip.")
    parser.add_argument("--xlsx", required=True, type=Path, help="Input .xlsx file path.")
    parser.add_argument("--sample-json", type=Path, default=DEFAULT_SAMPLE_JSON, help="Template sample JSON path.")
    parser.add_argument("--zip", required=True, type=Path, help="Output zip path.")
    parser.add_argument("--out-dir", required=True, type=Path, help="Directory for generated JSON files.")
    parser.add_argument("--sheet", help="Sheet name. Defaults to first sheet.")
    parser.add_argument("--start-index", type=int, default=1, help="1-based data row index, excluding header.")
    parser.add_argument("--count", type=int, help="Number of data rows to scan. Defaults to all rows.")
    parser.add_argument("--style", default="default", help="Style suffix for videoId.")
    parser.add_argument("--video-column", default="压缩后视频链接", help="Column used as styles.<style>.videoSrc.")
    parser.add_argument("--poster-column", default="视频poster", help="Column used as styles.<style>.videoPoster.")
    parser.add_argument("--clean", action="store_true", help="Remove out-dir before generation.")
    parser.add_argument("--continue-on-error", action="store_true", help="Skip invalid rows and still create zip.")
    return parser.parse_args(argv)


def main(argv: list[str]) -> int:
    try:
        result = make_package(parse_args(argv))
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if result["ok"] else 2
    except Exception as exc:  # noqa: BLE001 - CLI should return JSON error for workbench integration.
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False, indent=2), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
